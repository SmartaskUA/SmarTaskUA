# modules/excel_export.py
from __future__ import annotations

import os
import time
from pathlib import Path
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def export_schedule_to_excel(schedule_data: List[List[str]], output_path: str) -> str:
    """
    Escreve schedule_data (formato [header, *rows] tal como devolvido por
    build_output_rows()/solve()) para um ficheiro .xlsx em output_path.

    Retorna o caminho final onde o ficheiro foi gravado.
    """
    if not schedule_data:
        raise ValueError("schedule_data está vazio; nada para exportar.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule"

    for row in schedule_data:
        ws.append(row)

    # Cabeçalho a destacar
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Fixa a linha de cabeçalho e a coluna do employee_id ao fazer scroll
    ws.freeze_panes = "B2"

    # Largura de coluna aproximada ao conteúdo
    for col_idx, column_cells in enumerate(ws.columns, start=1):
        max_length = max(
            (len(str(c.value)) for c in column_cells if c.value is not None),
            default=8,
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_length + 2, 24)

    resolved_path = Path(output_path)
    resolved_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(resolved_path)

    return str(resolved_path)


def build_export_filename(algorithm_name: str, task_id: str) -> str:
    safe_algorithm_name = "".join(c if c.isalnum() or c in "-_" else "_" for c in algorithm_name)
    timestamp = time.strftime("%Y%m%dT%H%M%S")
    return f"{safe_algorithm_name}_{task_id}_{timestamp}.xlsx"