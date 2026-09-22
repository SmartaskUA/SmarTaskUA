"""
create_excel.py — Generates GA_Results.xlsx with clean comparison tables.

Sheet "2 Turnos": ILP / CSP / GA / GA+PP / Elite5 / Memético+PP (Vacation 1)
Sheet "3 Turnos": ILP / CSP / GA / GA+PP (Vacation 1)

Usage:
    python create_excel.py
"""

import csv, os
import numpy as np
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE    = os.path.dirname(__file__)
OUTPUT  = os.path.join(BASE, "GA_Results.xlsx")

# ── helpers ──────────────────────────────────────────────────────────────────

def read_ga(path):
    """(mean_min, mean_ideal, mean_time, n_runs) from ga_results.csv or results.csv."""
    if not os.path.exists(path):
        return None
    rows = []
    with open(path, newline="") as f:
        for r in csv.DictReader(f):
            rows.append(r)
    if not rows:
        return None
    mins   = [int(r["min_coverage_unmet"])     for r in rows]
    ideals = [float(r["ideal_coverage_unmet"]) for r in rows]
    times  = [float(r["elapsed_s"])            for r in rows]
    return (round(np.mean(mins), 1),
            round(np.mean(ideals), 1),
            round(np.mean(times), 0),
            len(rows))

def read_pp(path):
    """(mean_ideal_pp, mean_ls_time) from pp_results.csv."""
    if not os.path.exists(path):
        return None
    rows = []
    with open(path, newline="") as f:
        for r in csv.DictReader(f):
            rows.append(r)
    if not rows:
        return None
    ideals = [float(r["ideal_unmet_pp"]) for r in rows]
    times  = [float(r["elapsed_s"])      for r in rows]
    return (round(np.mean(ideals), 1),
            round(np.mean(times), 2))

NA = "n/a"

# ── reference data (ILP / CSP, Vacation 1) ───────────────────────────────────

ILP_2 = {   # (min, ideal, time_s)
    "2T":  (16,  126,    4.4),
    "4T":  ( 0,  488,   16.4),
    "8T":  ( 0,  976,   56.9),
    "16T": ( 0, 1952,  293.0),
    "32T": NA,
}
CSP_2 = {
    "2T":  (16,  127,   23.6),
    "4T":  ( 0,  488,  294.0),
    "8T":  ( 0,  976,  963.0),
    "16T": ( 0, 1956, 2705.0),
    "32T": NA,
}
ILP_3 = {
    "2T":  (103,  488,    74.0),
    "4T":  (215,  976,   546.0),
    "8T":  (412, 1952,  2700.0),
    "16T": NA,
    "32T": NA,
}
CSP_3 = {
    "2T":  (103,  488,   249.0),
    "4T":  (215,  976,  5403.0),
    "8T":  (412, 1952,  5413.0),
    "16T": NA,
    "32T": NA,
}

TEAMS_2 = ["2T", "4T", "8T", "16T", "32T"]
TEAMS_3 = ["2T", "4T", "8T", "16T", "32T"]

SC_2 = {
    "2T":  "SMARTASK_SIMPLE_2025",
    "4T":  "SMARTASK_4TEAMS_2025",
    "8T":  "SMARTASK_8TEAMS_2025",
    "16T": "SMARTASK_16TEAMS_2025",
    "32T": "SMARTASK_32TEAMS_2025",
}
SC_3 = {
    "2T":  "SMARTASK_3SHIFTS_2TEAMS_2025",
    "4T":  "SMARTASK_3SHIFTS_4TEAMS_2025",
    "8T":  "SMARTASK_3SHIFTS_8TEAMS_2025",
    "16T": "SMARTASK_3SHIFTS_16TEAMS_2025",
    "32T": "SMARTASK_3SHIFTS_32TEAMS_2025",
}

# ── styling helpers ───────────────────────────────────────────────────────────

def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def thin_border():
    s = Side(style="thin", color="BFBFBF")
    return Border(left=s, right=s, top=s, bottom=s)

def hdr_font(white=True, bold=True, size=10):
    return Font(name="Calibri", bold=bold, color="FFFFFF" if white else "1F1F1F", size=size)

def data_font(bold=False, size=10):
    return Font(name="Calibri", bold=bold, color="1F1F1F", size=size)

def center():
    return Alignment(horizontal="center", vertical="center", wrap_text=False)

def right():
    return Alignment(horizontal="right", vertical="center")

# Color palette
C = {
    "title":    "1F3864",   # dark navy
    "ilp":      "2E75B6",   # blue
    "csp":      "548235",   # green
    "ga":       "C55A11",   # orange
    "gapp":     "7030A0",   # purple
    "elite":    "833C00",   # brown
    "memetic":  "375623",   # dark green
    "sub":      "D6DCE4",   # light grey sub-header
    "row_alt":  "F2F2F2",   # alternating row
    "na":       "BFBFBF",   # grey for n/a cells
}

GROUP_COLORS = {
    "ILP":            C["ilp"],
    "CSP":            C["csp"],
    "GA":             C["ga"],
    "GA + PP":        C["gapp"],
    "GA Elite5":      C["elite"],
    "GA Memético+PP": C["memetic"],
}

# ── sheet builder ─────────────────────────────────────────────────────────────

def style_cell(cell, value, is_header=False, group_color=None,
               is_sub=False, is_na=False, alt_row=False, bold=False):
    cell.value = value
    cell.border = thin_border()
    cell.alignment = center() if (is_header or is_sub) else right()

    if is_header and group_color:
        cell.fill = fill(group_color)
        cell.font = hdr_font(white=True, bold=True)
    elif is_sub:
        cell.fill = fill(C["sub"])
        cell.font = hdr_font(white=False, bold=True, size=9)
    elif is_na:
        cell.fill = fill(C["na"])
        cell.font = data_font()
        cell.alignment = center()
    elif alt_row:
        cell.fill = fill(C["row_alt"])
        cell.font = data_font(bold=bold)
    else:
        cell.font = data_font(bold=bold)


def fmt_time(t):
    """Format seconds to readable string."""
    if t is None or t == NA:
        return NA
    t = float(t)
    if t < 60:
        return f"{t:.0f}s"
    m = int(t // 60)
    s = t % 60
    return f"{m}m {s:.0f}s"


def build_sheet(ws, title, teams, sc_map, results_dir, ilp_ref, csp_ref,
                results3=False, show_elite=False, show_memetic=False):
    """Build one comparison sheet."""

    # Column layout: [Equipas | ILP×3 | CSP×3 | GA×3 | GA+PP×2 | Elite5×3? | Mem+PP×2? ]
    groups = [
        ("ILP",            ["Mín.", "Ideal", "Tempo"]),
        ("CSP",            ["Mín.", "Ideal", "Tempo"]),
        ("GA",             ["Mín. (méd)", "Ideal (méd)", "Tempo (méd)"]),
        ("GA + PP",        ["Ideal (méd)", "Tempo total"]),
    ]
    if show_elite:
        groups.append(("GA Elite5",      ["Mín. (méd)", "Ideal (méd)", "Tempo (méd)"]))
    if show_memetic:
        groups.append(("GA Memético+PP", ["Ideal (méd)", "Tempo total"]))

    # Compute column positions
    col_start = 2  # column B
    group_ranges = []
    col = col_start
    for gname, subs in groups:
        group_ranges.append((gname, col, col + len(subs) - 1, subs))
        col += len(subs)
    total_cols = col - 1

    # ── Row 1: sheet title ────────────────────────────────────────────────────
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
    c = ws.cell(1, 1)
    c.value = title
    c.fill = fill(C["title"])
    c.font = Font(name="Calibri", bold=True, color="FFFFFF", size=13)
    c.alignment = center()
    ws.row_dimensions[1].height = 22

    # ── Row 2: group headers ──────────────────────────────────────────────────
    ws.merge_cells(start_row=2, start_column=1, end_row=3, end_column=1)
    c = ws.cell(2, 1)
    c.value = "Equipas"
    c.fill = fill(C["title"])
    c.font = hdr_font(white=True, bold=True)
    c.alignment = center()
    c.border = thin_border()

    for gname, c1, c2, subs in group_ranges:
        color = GROUP_COLORS[gname]
        ws.merge_cells(start_row=2, start_column=c1, end_row=2, end_column=c2)
        c = ws.cell(2, c1)
        style_cell(c, gname, is_header=True, group_color=color)

    # ── Row 3: sub-headers ────────────────────────────────────────────────────
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 18

    for gname, c1, c2, subs in group_ranges:
        for i, sub in enumerate(subs):
            c = ws.cell(3, c1 + i)
            style_cell(c, sub, is_sub=True)

    # ── Rows 4+: data ─────────────────────────────────────────────────────────
    for row_i, team in enumerate(teams):
        r = 4 + row_i
        alt = (row_i % 2 == 1)
        ws.row_dimensions[r].height = 16

        # Team label
        c = ws.cell(r, 1)
        style_cell(c, team, alt_row=alt, bold=True)
        c.alignment = center()

        col = col_start

        # ILP
        ilp = ilp_ref.get(team, NA)
        if ilp == NA:
            for _ in range(3):
                cc = ws.cell(r, col); style_cell(cc, NA, is_na=True); col += 1
        else:
            min_v, ideal_v, time_v = ilp
            for val in [min_v, ideal_v, fmt_time(time_v)]:
                cc = ws.cell(r, col); style_cell(cc, val, alt_row=alt); col += 1

        # CSP
        csp = csp_ref.get(team, NA)
        if csp == NA:
            for _ in range(3):
                cc = ws.cell(r, col); style_cell(cc, NA, is_na=True); col += 1
        else:
            min_v, ideal_v, time_v = csp
            for val in [min_v, ideal_v, fmt_time(time_v)]:
                cc = ws.cell(r, col); style_cell(cc, val, alt_row=alt); col += 1

        # GA (alone)
        sc = sc_map[team]
        if results3:
            ga_path = os.path.join(BASE, results_dir, sc, "results.csv")
        else:
            ga_path = os.path.join(BASE, results_dir, sc, "ga_results.csv")
        ga = read_ga(ga_path)
        if ga is None:
            for _ in range(3):
                cc = ws.cell(r, col); style_cell(cc, NA, is_na=True); col += 1
        else:
            mean_min, mean_ideal, mean_time, n = ga
            suffix = f" (n={n})" if n < 10 else ""
            for val in [mean_min, mean_ideal, fmt_time(mean_time) + suffix]:
                cc = ws.cell(r, col); style_cell(cc, val, alt_row=alt); col += 1

        # GA + PP
        if results3:
            pp_path = os.path.join(BASE, results_dir, sc, "pp_results.csv")
        else:
            pp_path = os.path.join(BASE, results_dir, sc, "pp_results.csv")
        pp = read_pp(pp_path)
        if pp is None:
            for _ in range(2):
                cc = ws.cell(r, col); style_cell(cc, NA, is_na=True); col += 1
        else:
            mean_ideal_pp, mean_ls_time = pp
            if ga:
                total_time = ga[2] + mean_ls_time
            else:
                total_time = None
            for val in [mean_ideal_pp, fmt_time(total_time) if total_time else NA]:
                cc = ws.cell(r, col); style_cell(cc, val, alt_row=alt); col += 1

        # GA Elite5 (2-shift only)
        if show_elite:
            e_ga_path = os.path.join(BASE, "results_elite5", sc, "ga_results.csv")
            e_ga = read_ga(e_ga_path)
            if e_ga is None:
                for _ in range(3):
                    cc = ws.cell(r, col); style_cell(cc, NA, is_na=True); col += 1
            else:
                mean_min, mean_ideal, mean_time, n = e_ga
                for val in [mean_min, mean_ideal, fmt_time(mean_time)]:
                    cc = ws.cell(r, col); style_cell(cc, val, alt_row=alt); col += 1

        # GA Memético + PP (2-shift only)
        if show_memetic:
            m_path  = os.path.join(BASE, "results_memetic_best", sc, "ma_results.csv")
            mp_path = os.path.join(BASE, "results_memetic_best", sc, "pp_results.csv")
            m_ga = read_ga(m_path)
            m_pp = read_pp(mp_path)
            if m_pp is None:
                for _ in range(2):
                    cc = ws.cell(r, col); style_cell(cc, NA, is_na=True); col += 1
            else:
                mean_ideal_pp, mean_ls_time = m_pp
                if m_ga:
                    total_time = m_ga[2] + mean_ls_time
                else:
                    total_time = None
                for val in [mean_ideal_pp, fmt_time(total_time) if total_time else NA]:
                    cc = ws.cell(r, col); style_cell(cc, val, alt_row=alt); col += 1

    # ── Column widths ─────────────────────────────────────────────────────────
    ws.column_dimensions["A"].width = 9
    for gname, c1, c2, subs in group_ranges:
        for i in range(c2 - c1 + 1):
            col_letter = get_column_letter(c1 + i)
            ws.column_dimensions[col_letter].width = 12

    # ── Notes row ─────────────────────────────────────────────────────────────
    note_row = 4 + len(teams) + 1
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=total_cols)
    c = ws.cell(note_row, 1)
    c.value = ("Vacation Scenario 1 (Sequential Rotation, 30 dias contínuos). "
               "GA e GA+PP: média de 10 runs. "
               "Mín. = mínimos não preenchidos; Ideal = ideais não preenchidos. "
               "ILP/CSP: run único. GA+PP tempo total inclui GA + local search.")
    c.font = Font(name="Calibri", italic=True, color="595959", size=9)
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[note_row].height = 28


# ── main ──────────────────────────────────────────────────────────────────────

def main():
    wb = openpyxl.Workbook()
    ws2 = wb.active
    ws2.title = "2 Turnos"
    ws3 = wb.create_sheet("3 Turnos")

    build_sheet(
        ws2,
        title="2 Turnos — GA vs ILP vs CSP (Vacation Scenario 1)",
        teams=TEAMS_2,
        sc_map=SC_2,
        results_dir="results_final",
        results3=False,
        ilp_ref=ILP_2,
        csp_ref=CSP_2,
        show_elite=True,
        show_memetic=True,
    )

    build_sheet(
        ws3,
        title="3 Turnos — GA vs ILP vs CSP (Vacation Scenario 1)",
        teams=TEAMS_3,
        sc_map=SC_3,
        results_dir="results_3shifts",
        results3=True,
        ilp_ref=ILP_3,
        csp_ref=CSP_3,
        show_elite=False,
        show_memetic=False,
    )

    wb.save(OUTPUT)
    print(f"Saved: {OUTPUT}")


if __name__ == "__main__":
    main()
